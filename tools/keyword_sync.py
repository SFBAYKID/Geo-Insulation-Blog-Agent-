"""Daily Drive workbook to Airtable sync, isolated from the read-only blog adapter."""

from __future__ import annotations

import argparse
import fcntl
import hashlib
import io
import json
import os
import time
from dataclasses import asdict
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

import requests

from tools.keyword_sync_model import FIELDS, build_plan, key, normalize, parse_rows


def save(path: Path, value: object) -> None:
    """Atomically persist private state before issuing any external write."""
    temp = path.with_suffix(path.suffix + ".tmp")
    with temp.open("w") as handle:
        os.chmod(temp, 0o600)
        json.dump(value, handle, indent=2)
        handle.flush()
        os.fsync(handle.fileno())
    temp.replace(path)


class Airtable:
    """Small rate-limited adapter; mutations are never blindly retried."""

    def __init__(self, config: dict[str, Any]) -> None:
        self.session = requests.Session()
        self.session.headers["Authorization"] = "Bearer " + config["AIRTABLE_TOKEN"]
        self.base = config["AIRTABLE_BASE_ID"]
        self.table = config["AIRTABLE_TABLE_ID"]

    def request(self, method: str, path: str, **kwargs: Any) -> dict[str, Any]:
        """Use sanitized failures and leave failed writes for operator inspection."""
        time.sleep(0.3)
        response = self.session.request(
            method, "https://api.airtable.com/v0/" + path, timeout=45, **kwargs
        )
        if not response.ok:
            raise RuntimeError(f"Airtable {method} failed (HTTP {response.status_code})")
        return dict(response.json())

    def schema(self) -> dict[str, str]:
        """Confirm all mapped fields exist and are writable scalar types."""
        tables = self.request("GET", f"meta/bases/{self.base}/tables")["tables"]
        table = next(t for t in tables if t["id"] == self.table)
        types = {f["name"]: f["type"] for f in table["fields"]}
        allowed = {
            "singleLineText",
            "multilineText",
            "singleSelect",
            "number",
            "url",
            "currency",
            "percent",
        }
        if any(name not in types or types[name] not in allowed for name in FIELDS):
            raise ValueError("Airtable mapped fields are missing or their types changed")
        return {name: types[name] for name in FIELDS}

    def records(self) -> list[dict[str, Any]]:
        """Read the entire table, including records hidden from the blog's view."""
        rows: list[dict[str, Any]] = []
        params: dict[str, Any] = {"pageSize": 100}
        while True:
            page = self.request("GET", f"{self.base}/{self.table}", params=params)
            rows.extend(page["records"])
            if not page.get("offset"):
                return rows
            params["offset"] = page["offset"]


def read_source(config: dict[str, Any]) -> tuple[list[dict[str, Any]], dict[str, Any]]:
    """Read the exact configured XLSX file and reject formulas or concurrent edits."""
    from google.auth.transport.requests import AuthorizedSession
    from google.oauth2 import service_account
    from openpyxl import load_workbook

    credentials = service_account.Credentials.from_service_account_file(
        config["credentials_path"], scopes=["https://www.googleapis.com/auth/drive.readonly"]
    )
    session = AuthorizedSession(credentials)
    url = "https://www.googleapis.com/drive/v3/files/" + config["source_file_id"]
    params = {"supportsAllDrives": "true", "fields": "id,name,mimeType,version,modifiedTime,size"}
    response = session.get(url, params=params, timeout=45)
    if not response.ok:
        raise RuntimeError(f"Source access failed (HTTP {response.status_code})")
    metadata = response.json()
    if metadata["mimeType"] != "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet":
        raise ValueError("Source file format changed; expected XLSX")
    if int(metadata.get("size", 0)) > 20_000_000:
        raise ValueError("Source file exceeds 20 MB limit")
    response = session.get(url, params={"supportsAllDrives": "true", "alt": "media"}, timeout=60)
    if not response.ok:
        raise RuntimeError(f"Source download failed (HTTP {response.status_code})")
    content = response.content
    if len(content) > 20_000_000:
        raise ValueError("Source download exceeds 20 MB limit")
    check = session.get(url, params=params, timeout=45)
    if not check.ok or check.json().get("version") != metadata.get("version"):
        raise ValueError("Source changed during download; try the next check")
    workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=False)
    try:
        sheet = workbook[config["sheet_name"]]
        if sheet.max_row > 20000 or sheet.max_column > 100:
            raise ValueError("Unexpected workbook dimensions")
        rows: list[list[object]] = []
        for row in sheet.iter_rows(max_col=10):
            if any(cell.data_type in {"f", "e"} for cell in row):
                raise ValueError("Source contains formulas or errors; mapping needs review")
            rows.append([cell.value for cell in row])
    finally:
        workbook.close()
    metadata["sha256"] = hashlib.sha256(content).hexdigest()
    return parse_rows(rows), metadata


def run(config: dict[str, Any], apply: bool) -> dict[str, Any]:
    """Plan, optionally apply, verify, and checkpoint one complete sync under a lock."""
    os.umask(0o077)
    directory = Path(config["state_dir"])
    directory.mkdir(parents=True, exist_ok=True, mode=0o700)
    with (directory / "sync.lock").open("a") as lock:
        fcntl.flock(lock, fcntl.LOCK_EX | fcntl.LOCK_NB)
        return sync_locked(config, directory, apply)


def sync_locked(config: dict[str, Any], directory: Path, apply: bool) -> dict[str, Any]:
    """Keep an unresolved write journal after any failure; never retry ambiguous creates."""
    pending = directory / "pending.json"
    if pending.exists():
        raise RuntimeError("Unresolved write journal: inspect pending.json before further sync")
    state_path = directory / "state.json"
    state = json.loads(state_path.read_text()) if state_path.exists() else {}
    binding = {k: config[k] for k in ("source_file_id", "sheet_name")}
    air_config = json.loads(Path(config["airtable_config_path"]).read_text())
    binding.update({k: air_config[k] for k in ("AIRTABLE_BASE_ID", "AIRTABLE_TABLE_ID")})
    if state and state["binding"] != binding:
        raise ValueError("Saved state belongs to a different source or target")
    source, metadata = read_source(config)
    api = Airtable(air_config)
    types = api.schema()
    records = api.records()
    previous = state.get("source", {})
    plan = build_plan(source, records, previous, types)
    stamp = datetime.now(timezone.utc).isoformat()
    report = {
        "checked_at": stamp,
        "source_file": metadata,
        "apply": apply,
        "source_rows": len(source),
        "target_rows": len(records),
        **asdict(plan),
    }
    save(directory / "last-plan.json", report)
    if plan.conflicts:
        raise ValueError(f"Sync held: {len(plan.conflicts)} conflict(s); see last-plan.json")
    summary = {
        "source_rows": len(source),
        "created": len(plan.creates),
        "updated": len(plan.updates),
        "unchanged": plan.unchanged,
        "missing_from_source_retained": len(plan.missing),
        "applied": apply,
    }
    if not apply:
        return summary
    # A durable full backup and journal precede writes, including for crash recovery.
    if plan.creates or plan.updates:
        save(directory / ("backup-" + str(time.time_ns()) + ".json"), records)
        save(pending, report)
        # Recheck all records before starting, reducing the window for manual edits.
        if api.records() != records:
            raise ValueError("Airtable changed during planning; inspect pending journal")
    expected = {record["id"]: dict(record["fields"]) for record in records}
    for update in plan.updates:
        result = api.request(
            "PATCH",
            f"{api.base}/{api.table}/{update['id']}",
            json={"fields": update["fields"], "typecast": True},
        )
        if result.get("id") != update["id"]:
            raise RuntimeError("Update response identity mismatch")
        expected[update["id"]].update(update["fields"])
    for fields in plan.creates:
        result = api.request(
            "POST",
            f"{api.base}/{api.table}",
            json={"records": [{"fields": fields}], "typecast": True},
        )
        created = result["records"]
        if len(created) != 1:
            raise RuntimeError("Unexpected creation response")
        expected[created[0]["id"]] = fields
    actual = api.records() if plan.creates or plan.updates else records
    actual_by_id = {r["id"]: r["fields"] for r in actual}
    if len(actual) != len(records) + len(plan.creates):
        raise RuntimeError("Record count changed unexpectedly; inspect pending journal")
    for record_id, fields in expected.items():
        if record_id not in actual_by_id or normalize(actual_by_id[record_id], types) != normalize(
            fields, types
        ):
            raise RuntimeError("Airtable verification mismatch; inspect pending journal")
    identities = [key(r["fields"]) for r in actual if r["fields"].get("Keyword")]
    if len(identities) != len(set(identities)):
        raise RuntimeError("Duplicate identity after sync; inspect pending journal")
    # Retain missing historical keys so a subsequent rename is never treated as new.
    saved_source = dict(previous)
    saved_source.update({key(row): normalize(row, types) for row in source})
    save(
        state_path,
        {"binding": binding, "source": saved_source, "checked_at": stamp, "source_file": metadata},
    )
    if pending.exists():
        pending.rename(directory / ("applied-" + str(time.time_ns()) + ".json"))
    save(directory / "last-result.json", {"checked_at": stamp, **summary})
    return summary


def main() -> None:
    """Run a read-only preview by default; --apply enables the authorized sync."""
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--config", required=True)
    parser.add_argument("--apply", action="store_true")
    args = parser.parse_args()
    try:
        config = json.loads(Path(args.config).read_text())
        print(json.dumps(run(config, args.apply)))
    except Exception as error:
        # Provider errors can contain request details; do not dump a traceback or tokens.
        safe = str(error) if isinstance(error, (ValueError, RuntimeError)) else type(error).__name__
        raise SystemExit("Keyword sync failed: " + safe) from None


if __name__ == "__main__":
    main()
